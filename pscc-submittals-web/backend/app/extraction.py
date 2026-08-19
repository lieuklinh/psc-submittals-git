"""
extraction.py — everything that talks to the existing extraction FastAPI
backend (Alekhya's psc-submittals-git repo, EXTRACTION_BACKEND_URL), OpenAI,
and the PDF itself. This is a direct port of the equivalent functions from
pscc-submittals-agent/app.py (the Streamlit version) — same prompts, same
heuristics — just relocated so the new FastAPI backend can call them
without depending on Streamlit.
"""

import base64
import io
import json
import re
import zipfile

import fitz  # PyMuPDF
import requests
from openai import OpenAI

from . import config

try:
    import pymupdf4llm
    HAS_PYMUPDF4LLM = True
except ImportError:
    HAS_PYMUPDF4LLM = False


def _raise_with_response_detail(resp, endpoint_name):
    if resp.ok:
        return
    try:
        detail = resp.json()
    except Exception:
        detail = resp.text[:500]
    raise RuntimeError(f"{endpoint_name} returned {resp.status_code}: {detail}")


def call_extract_sections(filename: str, b64_content: str) -> bytes:
    resp = requests.post(
        f"{config.EXTRACTION_BACKEND_URL}/extract-submittals-sections",
        json={"filename": filename, "file_content": b64_content}, timeout=600,
    )
    _raise_with_response_detail(resp, "extract-submittals-sections")
    return resp.content


def call_extract_log(filename: str, b64_content: str) -> bytes:
    resp = requests.post(
        f"{config.EXTRACTION_BACKEND_URL}/extract-submittals-log",
        json={"filename": filename, "file_content": b64_content}, timeout=600,
    )
    _raise_with_response_detail(resp, "extract-submittals-log")
    return resp.content


def call_create_structure(filename: str, b64_content: str, project_info: dict) -> bytes:
    resp = requests.post(
        f"{config.EXTRACTION_BACKEND_URL}/create-submittal-structure",
        json={"filename": filename, "file_content": b64_content, "project_info": project_info},
        timeout=600,
    )
    _raise_with_response_detail(resp, "create-submittal-structure")
    return resp.content


def call_build_index(filename: str, b64_content: str):
    if not config.AGENT_URL:
        return None
    try:
        resp = requests.post(
            f"{config.AGENT_URL}/build-index",
            json={"filename": filename, "file_content": b64_content},
            timeout=1200,
        )
        if resp.status_code == 503:
            return None
        resp.raise_for_status()
        return resp.json().get("doc_id")
    except Exception:
        return None


def call_agent_chat(message: str, doc_id: str | None) -> str:
    if not config.AGENT_URL or not doc_id:
        return (
            "The Q&A agent isn't connected yet. Once the agent API is running and "
            "AGENT_URL is set, your questions will be answered from the spec-book index.\n\n"
            f"You asked: {message}"
        )
    try:
        resp = requests.post(
            f"{config.AGENT_URL}/chat",
            json={"message": message, "doc_id": doc_id},
            timeout=180,
        )
        if resp.status_code == 503:
            detail = ""
            try:
                detail = resp.json().get("detail", "")
            except Exception:
                pass
            return f"Agent not ready ({detail or 'not loaded'}).\n\nYou asked: {message}"
        resp.raise_for_status()
        return resp.json().get("answer", "(the agent returned no answer)")
    except Exception as e:
        return f"Couldn't reach the Q&A agent ({e}).\n\nYou asked: {message}"


def scan_pdf_for_irregular_codes(pdf_bytes: bytes, n_pages: int = 25):
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    pages_to_scan = min(n_pages, len(doc))
    lines = []
    for i in range(pages_to_scan):
        lines.extend(doc[i].get_text("text").split("\n"))
    doc.close()

    standard = re.compile(r"^\d{2}\s\d{2}\s\d{2}$")
    irregular_patterns = [
        ("attached suffix letters", re.compile(r"^\d{2}\s\d{2}\s\d{2}[A-Za-z]+$")),
        ("dot-separated", re.compile(r"^\d{2}\.\d{2}\.\d{2}$")),
        ("decimal sub-level", re.compile(r"^\d{2}\s\d{2}\.\d{2}$")),
        ("non-CSI custom prefix", re.compile(r"^[A-Za-z]{2,6}-\d+[A-Za-z]?$")),
        ("inline note/parenthetical", re.compile(r"^\d{2}\s\d{2}\s\d{2}\s*\(.*\)$")),
    ]

    flagged = []
    for idx, line in enumerate(lines):
        stripped = line.strip()
        if not stripped or standard.match(stripped):
            continue
        for reason, pattern in irregular_patterns:
            if pattern.match(stripped):
                context = next((l.strip() for l in lines[idx + 1:idx + 3] if l.strip()), "")
                flagged.append({"code": stripped, "reason": reason, "context": context})
                break
    return flagged


def extract_first_n_pages(pdf_bytes, n_pages=10, mode="Plain text", char_limit=12000):
    if mode.startswith("Markdown") and HAS_PYMUPDF4LLM:
        import tempfile
        import os
        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
            tmp.write(pdf_bytes)
            tmp_path = tmp.name
        try:
            doc = fitz.open(tmp_path)
            pages_to_use = min(n_pages, len(doc))
            doc.close()
            md_text = pymupdf4llm.to_markdown(tmp_path, pages=list(range(pages_to_use)))
            return md_text[:char_limit], pages_to_use
        finally:
            os.unlink(tmp_path)

    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    pages_to_use = min(n_pages, len(doc))
    text = ""
    for page_num in range(pages_to_use):
        text += f"\n--- Page {page_num + 1} ---\n{doc[page_num].get_text('text', sort=True)}"
    doc.close()
    return text[:char_limit], pages_to_use


PROJECT_INFO_SCHEMA_KEYS = [
    "project_name", "ccua_project_number", "pscc_job_number",
    "engineer_name", "engineer_address", "contractor_name", "contractor_address",
    "owner_name", "owner_address", "prepared_by",
]

PROJECT_INFO_PROMPT_TEMPLATE = """Extract the following project info (even if wording varies or formatting is inconsistent):

- Project Name (also appears as "Project Title", "Project", or on title page)
- Project Number (also called "Project No.", "Contract No.", "CCUA Project Number", "Project ID")
- PSCC Job Number (also appears as "PSCC Project Number", "Job No.", "Job Number")
- Engineer Name and Address (may appear as a firm name, contact block, or footer stamp)
- Contractor Name and Address (may appear on signature page or contract information sheet)
- Owner Name and Address (may appear as "Owner", "Client", "Utility Authority", etc.)
- Prepared By (may be an engineer, consultant, or preparation firm listed on cover/title page)

IMPORTANT:
- If a value appears multiple times, use the most complete version.
- If a field is partially available, populate what is available and leave the rest empty.
- If a field is truly not present, return an empty string "" for that key.
- Never omit any key.

Return ONLY strict JSON with keys: project_name, ccua_project_number, pscc_job_number, engineer_name, engineer_address, contractor_name, contractor_address, owner_name, owner_address, prepared_by

PDF Text (first {pages_to_use} pages):
\"\"\"
{pdf_text}
\"\"\"
"""


def extract_project_info_via_openai(pdf_bytes, api_key, model, mode="Plain text"):
    if not api_key:
        raise ValueError("No OpenAI API key configured on the server.")

    pdf_text, pages_used = extract_first_n_pages(pdf_bytes, n_pages=10, mode=mode)
    prompt = PROJECT_INFO_PROMPT_TEMPLATE.format(pages_to_use=pages_used, pdf_text=pdf_text)

    client = OpenAI(api_key=api_key)
    response = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": "You are a helpful assistant that extracts project information from construction PDFs. Always return valid JSON only."},
            {"role": "user", "content": prompt},
        ],
        temperature=0,
        max_tokens=500,
    )

    raw = response.choices[0].message.content.strip()
    if raw.startswith("```json"):
        raw = raw.replace("```json", "").replace("```", "").strip()
    elif raw.startswith("```"):
        raw = raw.replace("```", "").strip()

    try:
        info = json.loads(raw)
    except json.JSONDecodeError:
        info = {}

    for key in PROJECT_INFO_SCHEMA_KEYS:
        info.setdefault(key, "")
    return info, pages_used


def detect_section_code_anomalies(sections):
    flagged = []
    for s in sections:
        desc = (s.get("description") or "").strip()
        code = (s.get("section_code") or "").strip()
        if " - " not in desc:
            continue
        embedded_code = desc.split(" - ", 1)[0].strip()
        if embedded_code and embedded_code != code:
            flagged.append({**s, "embedded_code": embedded_code})
    return flagged


def explain_anomaly_via_openai(original_code, embedded_code, description, api_key, model):
    if not api_key:
        return "OpenAI key not configured — pick a value manually.", None

    prompt = f"""A construction submittal section list has a mismatch:
- Extracted section code: "{original_code}"
- Code embedded in that same row's description text: "{embedded_code}"
- Full description: "{description}"

In 1-2 plain sentences, explain the likely reason for this mismatch (e.g. a
state-specific suffix, a sub-section number, an extraction quirk). Then
suggest which value should be used going forward.

Return ONLY strict JSON: {{"explanation": "...", "suggested_code": "..."}}"""

    try:
        client = OpenAI(api_key=api_key)
        response = client.chat.completions.create(
            model=model,
            messages=[{"role": "user", "content": prompt}],
            temperature=0,
            max_tokens=200,
        )
        raw = response.choices[0].message.content.strip()
        if raw.startswith("```"):
            raw = raw.strip("`").replace("json", "", 1).strip()
        data = json.loads(raw)
        return data.get("explanation", ""), data.get("suggested_code")
    except Exception as e:
        return f"Couldn't get an explanation ({e}) — pick a value manually.", None


def parse_log_for_workspace(log_bytes: bytes):
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(log_bytes), data_only=True)
    ws = wb["Submittal Resp"] if "Submittal Resp" in wb.sheetnames else wb.active

    header_row = None
    col_map = {}
    FIELD_LABELS = {
        "section": "section_code",
        "description": "description",
        "type": "submittal_type",
        "sub/vendor": "sub_vendor",
        "due": "due_date",
        "float": "float_days",
    }

    for row in ws.iter_rows(min_row=1, max_row=min(10, ws.max_row)):
        found_this_row = {}
        for cell in row:
            if not isinstance(cell.value, str):
                continue
            label = cell.value.strip().lower()
            for prefix, field in FIELD_LABELS.items():
                if label.startswith(prefix):
                    found_this_row[field] = cell.column
            if label.startswith("recv"):
                found_this_row["received_date"] = cell.column
        if "section_code" in found_this_row and "description" in found_this_row:
            header_row = row[0].row
            col_map = found_this_row
            break

    if not header_row:
        return []

    sections = []
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        code_cell = row[col_map["section_code"] - 1].value
        if not code_cell or not str(code_cell).strip():
            continue

        def _cell(field):
            col = col_map.get(field)
            if not col:
                return None
            val = row[col - 1].value
            return str(val).strip() if val is not None else None

        sections.append(
            {
                "section_code": str(code_cell).strip(),
                "description": _cell("description") or "",
                "submittal_type": _cell("submittal_type"),
                "sub_vendor": _cell("sub_vendor"),
                "due_date": _cell("due_date"),
                "received_date": _cell("received_date"),
                "float_days": _cell("float_days"),
            }
        )
    return sections


def bundle_outputs_into_zip(sections_bytes, log_bytes, structure_zip_bytes, base_name):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as out_zip:
        out_zip.writestr(f"{base_name}_submittal_sections.xlsx", sections_bytes)
        out_zip.writestr(f"{base_name}_submittals_log.xlsx", log_bytes)
        with zipfile.ZipFile(io.BytesIO(structure_zip_bytes), "r") as struct_zip:
            for item in struct_zip.infolist():
                data = struct_zip.read(item.filename)
                out_zip.writestr(f"Submittal_Structure/{item.filename}", data)
    buffer.seek(0)
    return buffer.getvalue()


def extract_section_full_text(pdf_bytes, sections: list[dict]):
    """Returns {specification_section_id: chunk_text} by locating each
    section's section_number in the raw PDF text layer. Heuristic — same
    approach as the Streamlit version; sections that don't match verbatim
    simply get no text (and thus no AI summary input)."""
    if not sections:
        return {}

    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    full_text = "\n".join(page.get_text("text") for page in doc)
    doc.close()

    positions = []
    for s in sections:
        idx = full_text.find(s["section_number"])
        if idx != -1:
            positions.append((idx, s))
    positions.sort(key=lambda p: p[0])

    result = {}
    for i, (idx, s) in enumerate(positions):
        end = positions[i + 1][0] if i + 1 < len(positions) else min(idx + 8000, len(full_text))
        chunk = full_text[idx:end].strip()[:6000]
        if chunk:
            result[s["id"]] = chunk
    return result


SECTION_SUMMARY_PROMPT = """List exactly 3 short bullet points summarizing the following construction specification section. Each bullet must be a short phrase of 4 to 10 words only — a fragment, not a full sentence — capturing one key idea, in this order:
1. The main scope of work
2. The major materials or systems involved
3. The principal submittal requirements

Example style (do not copy content, only match this brevity and phrasing style):
- preparation, placement, and finishing of concrete surfaces
- concrete mixes and aggregates
- detailed mix designs, shop drawings, and certifications

Start each line with "- ". Use only information from the provided specification text. Do not invent requirements. Return ONLY the 3 bullet lines — no heading, no intro, no closing sentence.

Specification section text:
\"\"\"
{text}
\"\"\"
"""


def generate_section_summary(section: dict, api_key: str, model: str):
    source_text = (section.get("section_text_reference") or "").strip()
    if not source_text:
        source_text = (section.get("section_name") or "").strip()
    if not source_text or not api_key:
        return None
    try:
        client = OpenAI(api_key=api_key)
        response = client.chat.completions.create(
            model=model,
            messages=[{"role": "user", "content": SECTION_SUMMARY_PROMPT.format(text=source_text[:6000])}],
            temperature=0.2,
            max_tokens=100,
        )
        return response.choices[0].message.content.strip()
    except Exception:
        return None
